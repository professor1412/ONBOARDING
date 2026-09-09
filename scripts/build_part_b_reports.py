"""Compact Part B memo and presentation exports.

Values are passed from the executed Part B notebook; this module formats them
and does not recalculate the airport analysis.
"""
from pathlib import Path
from xml.sax.saxutils import escape


def create_part_b_reports(summary, tables, figures, deliverables):
    deliverables = Path(deliverables)
    deliverables.mkdir(parents=True, exist_ok=True)
    critical = '21:00–04:00'
    decision = 'Do not fund a blanket airport-catchment signup bonus yet. Run a timed supply and return-loop pilot first.'
    findings = [
        ('B1 · The shortage is real, concentrated and repeatable.',
         f"Airport-terminal fulfilment is {summary['critical_fulfilment']:.1%} in {critical}, leaving "
         f"{summary['critical_unfulfilled']:,} requests unfulfilled over May–June ({summary['critical_unfulfilled_per_day']:.0f}/day, "
         f"{summary['critical_unfulfilled_share']:.1%} of the terminal total). Daytime fulfilment is "
         f"{summary['daytime_fulfilment']:.1%}. The gap is stable: May {summary['may_critical_fulfilment']:.1%}; "
         f"June {summary['june_critical_fulfilment']:.1%}."),
        ('B2 · The airport trip often strands supply after the ride.',
         f"In the 24-hour-buffer sample, suburban trips cancel at {summary['suburban_cancel_rate']:.1%} vs "
         f"{summary['core_cancel_rate']:.1%} for city-core (standardized gap {summary['standardized_cancel_gap_pp']:.1f} pp). "
         f"Only {summary['suburban_next_fare_rate']:.1%} of suburban trips get a return fare within 20 minutes vs "
         f"{summary['core_next_fare_rate']:.1%} for core trips. The file is a sample, has no unique captain ID or "
         "completion timestamp, and cannot establish captain earnings."),
        ('B3 · Acquisition is the wrong first intervention.',
         f"A 10% recovery of critical unfulfilled demand is {summary['pilot_extra_fulfilled_per_night']:.0f} additional "
         f"fulfilled requests/night. A 20-slot timed roster has illustrative capacity for {summary['pilot_capacity_at_assumed_productivity']:.0f} "
         f"completions/night; an INR {summary['pilot_expected_nightly_variable_cost']:,.0f}/night variable-cost scenario "
         f"is about INR {summary['pilot_cost_per_incremental_fulfilment_at_target']:.0f} per incremental fulfilment at the target. "
         "These are planning assumptions, not a permanent-program business case."),
    ]
    actions = [
        ('1. Run a 28-night timed availability pilot',
         'Roster 20 availability slots across the 21:00–04:00 window; randomize comparable terminal-hours to pilot/control. '
         'Measure incremental fulfilled requests per online captain-hour, unfulfilled demand, ETA and surge. Keep attendance and incentive logs.'),
        ('2. Fix the return loop before buying reach',
         'Test airport queue/dispatch improvements, return-fare matching and a bounded deadhead fallback. Use suburban/core '
         'return-fare and cancellation rates as guardrails, while measuring complaints, retention and actual captain earnings.'),
        ('3. Revisit targeted acquisition only after the pilot',
         'If timed supply improves demand fulfilment and utilization at an acceptable cost, run a small randomized catchment '
         'acquisition/bonus test with unique-captain IDs and a measured retention window. Do not extrapolate sampled trips into signups today.'),
    ]
    assumptions = (
        'Planning assumptions: 20 roster slots, 80% attendance, 0.75 completions per effective captain-hour, '
        'INR 30/hour availability incentive, INR 60 fallback capped at 20 captains/night. Costs exclude fixed engineering, '
        'airport fees, acquisition bonus and captain earnings. Hourly and trip extracts have different denominators.'
    )
    limits = (
        f"Data follow-ups: {summary['sample_exceeds_requests_cells']} terminal-hours have sampled trip rows above marketplace request counts; "
        f"{summary['cancelled_with_return']} cancelled records carry a return-fare flag; and 7,740 zone-hours have exactly one unfulfilled request. "
        'Validate the source definitions before any permanent spend.'
    )
    md = ['# Part B — Airport supply decision memo',
          'For Regional Leadership · Extract: 30 June 2026, 23:59 IST', decision, '## Findings']
    for title, text in findings: md += [f'**{title}** {text}']
    md += ['## Ranked actions']
    for title, text in actions: md += [f'### {title}', text]
    md += [assumptions, limits, 'Working: `analysis_partB.ipynb`; tables and figures: `outputs/part_b/`.']
    (deliverables / 'part_b_memo.md').write_text('\n\n'.join(md), encoding='utf-8')

    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    font_root = Path('/usr/share/fonts/truetype/dejavu')
    font, bold = 'Helvetica', 'Helvetica-Bold'
    if (font_root / 'DejaVuSans.ttf').exists():
        pdfmetrics.registerFont(TTFont('PartBSans', str(font_root / 'DejaVuSans.ttf')))
        pdfmetrics.registerFont(TTFont('PartBSansBold', str(font_root / 'DejaVuSans-Bold.ttf')))
        font, bold = 'PartBSans', 'PartBSansBold'
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name='BTitle', fontName=bold, fontSize=19, leading=23,
                              textColor=colors.HexColor('#193B50'), spaceAfter=9))
    styles.add(ParagraphStyle(name='BBody', fontName=font, fontSize=9.2, leading=13.1, spaceAfter=8))
    styles.add(ParagraphStyle(name='BSmall', fontName=font, fontSize=8.3, leading=11, spaceAfter=7))
    styles.add(ParagraphStyle(name='BHead', fontName=bold, fontSize=11, leading=14,
                              textColor=colors.HexColor('#245B78'), spaceBefore=5, spaceAfter=5))
    def p(text, style='BBody'): return Paragraph(escape(text), styles[style])
    memo = deliverables / 'part_b_memo.pdf'
    doc = SimpleDocTemplate(str(memo), pagesize=A4, rightMargin=40, leftMargin=40, topMargin=35, bottomMargin=35,
                            title='Part B — Airport supply decision memo', author='Supply analytics')
    flow = [p('Airport supply: solve the return loop before buying reach', 'BTitle'),
            p('DECISION MEMO  |  PART B  |  30 JUNE 2026 EXTRACT', 'BSmall'), p(decision), Spacer(1,5)]
    for title, text in findings: flow += [p(title,'BHead'),p(text)]
    flow += [p('What changes the decision','BHead'),p('A timed pilot that improves fulfilled demand and utilization at acceptable cost can justify a targeted acquisition test. '
        'The current files do not support a signup-bonus forecast.'),p(limits,'BSmall'),PageBreak(),p('Three actions, ranked','BTitle')]
    for title, text in actions: flow += [p(title,'BHead'),p(text)]
    flow += [p('Pilot working','BHead'),p(f"A 10% recovery target is {summary['pilot_extra_fulfilled_per_night']:.0f} extra fulfilled requests/night "
        f"and {summary['pilot_extra_per_30_day_month']:.0f} per 30-day month. Illustrative variable cost is "
        f"INR {summary['pilot_expected_nightly_variable_cost']:,.0f}/night (cap INR {summary['pilot_max_nightly_variable_cost']:,.0f}), "
        f"or INR {summary['pilot_cost_per_incremental_fulfilment_at_target']:.0f} per incremental fulfilment at target."),
        p(assumptions,'BSmall'),p('Audit trail: analysis_partB.ipynb and outputs/part_b/*.csv.','BSmall')]
    def footer(canvas, document):
        canvas.saveState(); canvas.setFont(font,7); canvas.setFillColor(colors.HexColor('#667680'))
        canvas.drawString(40,21,'SUPPLY ANALYTICS  /  PART B  /  Sample and planning assumptions are explicit')
        canvas.drawRightString(A4[0]-40,21,str(document.page)); canvas.restoreState()
    doc.build(flow, onFirstPage=footer, onLaterPages=footer)

    # A compact four-slide PDF and editable PPTX.
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from PIL import Image
    prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
    navy='193B50'; blue='245B78'; orange='DF803C'; pale='EDF3F6'; grey='5B6B76'
    def box(sl,x,y,w,h,fill='FFFFFF'):
        sh=sl.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h)); sh.fill.solid(); sh.fill.fore_color.rgb=RGBColor.from_string(fill); sh.line.fill.background(); return sh
    def text(sl,x,y,w,h,value,size=18,color=navy,bold=False):
        sh=sl.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h)); tf=sh.text_frame; tf.word_wrap=True; tf.margin_left=Inches(.02);tf.margin_right=Inches(.02)
        for i,line in enumerate(value.split('\n')):
            para=tf.paragraphs[0] if i==0 else tf.add_paragraph(); para.text=line; para.font.name='DejaVu Sans'; para.font.size=Pt(size); para.font.bold=bold; para.font.color.rgb=RGBColor.from_string(color); para.space_after=Pt(8)
        return sh
    def slide(title, sub):
        sl=prs.slides.add_slide(prs.slide_layouts[6]); box(sl,0,0,13.333,.12,blue); text(sl,.55,.35,12.2,.7,title,27,bold=True); text(sl,.58,1.05,12.1,.65,sub,13,color=grey); idx=len(prs.slides); text(sl,.58,7.12,11.8,.22,'PART B  |  Airport supply  |  Extract: 30 Jun 2026, 23:59 IST',9,color=grey); text(sl,12.3,7.10,.5,.28,f'{idx}/4',10,color=grey); return sl
    def picture(sl,path,x,y,w,h):
        with Image.open(path) as im: iw,ih=im.size
        scale=min(w/iw,h/ih); fw,fh=iw*scale,ih*scale; sl.shapes.add_picture(str(path),Inches(x+(w-fw)/2),Inches(y+(h-fh)/2),width=Inches(fw),height=Inches(fh))
    def card(sl,x,y,w,h,title,body): box(sl,x,y,w,h,pale); text(sl,x+.18,y+.13,w-.36,.55,title,20,bold=True); text(sl,x+.18,y+.82,w-.36,h-.95,body,16)

    sl=slide('Airport shortage is real; acquisition is not the first fix', 'Decision: run a timed supply and return-loop pilot before a catchment signup bonus.')
    for x,val,label in [(0.6,f"{summary['critical_fulfilment']:.1%}",'Critical-window fulfilment'),(4.85,f"{summary['critical_unfulfilled_per_day']:.0f}",'Unfulfilled/day'),(9.1,f"{summary['suburban_next_fare_rate']:.1%}",'Suburban return fare')]:
        box(sl,x,1.95,3.62,1.65,pale); text(sl,x+.18,2.1,3.25,.67,val,34,bold=True); text(sl,x+.18,2.85,3.25,.45,label,16)
    text(sl,.65,4.0,12,2.3,f"1   21:00–04:00 creates {summary['critical_unfulfilled_share']:.1%} of terminal unfulfilled demand.\n2   Suburban trips have {summary['standardized_cancel_gap_pp']:.1f} pp higher standardized cancellation.\n3   Test timed availability and return matching before buying reach.",22); sl.notes_slide.notes_text_frame.text=decision+' '+assumptions
    sl=slide('The gap is concentrated in the 21:00–04:00 window', 'May discovery, June validation; fulfilment uses request-weighted totals.')
    picture(sl,figures['airport_period_mismatch'],.55,1.75,12.2,4.75); text(sl,.65,6.55,12,.35,f"Critical fulfilment {summary['critical_fulfilment']:.1%} vs daytime {summary['daytime_fulfilment']:.1%}; online supply {summary['critical_mean_online_per_terminal_hour']:.1f} vs {summary['daytime_mean_online_per_terminal_hour']:.1f} per terminal-hour.",13,color=grey); sl.notes_slide.notes_text_frame.text=findings[0][1]
    sl=slide('Post-trip economics point to a return-loop problem', 'Airport-origin trip sample; primary return comparison excludes the final 24 hours.')
    picture(sl,figures['airport_trip_outcomes'],.55,1.8,7.0,4.7); card(sl,7.8,1.95,4.95,4.35,'Suburban vs city-core',f"Cancellation: {summary['suburban_cancel_rate']:.1%} vs {summary['core_cancel_rate']:.1%}.\n\nReturn fare within 20 min: {summary['suburban_next_fare_rate']:.1%} vs {summary['core_next_fare_rate']:.1%}.\n\nThe sample has no unique captain ID, completion timestamp or earnings."); sl.notes_slide.notes_text_frame.text=findings[1][1]
    sl=slide('Run the pilot, then decide whether to acquire', 'A controlled test gives Supply a measurable gate for permanent spend.')
    card(sl,.6,1.95,5.95,4.55,'Timed supply pilot',f"20 roster slots; 80% attendance.\n\nTarget 10% recovery: ~{summary['pilot_extra_fulfilled_per_night']:.0f} extra fulfilments/night.\n\nIllustrative cost: INR {summary['pilot_expected_nightly_variable_cost']:,.0f}/night; ~INR {summary['pilot_cost_per_incremental_fulfilment_at_target']:.0f}/incremental fulfilment at target.\n\nRandomize 28 nights; measure fulfilment/hour, ETA, surge and return fare.")
    card(sl,6.8,1.95,5.95,4.55,'Acquisition gate',f"No blanket catchment bonus today.\n\nProceed only if the pilot improves incremental fulfilment and return economics with acceptable complaints and captain retention.\n\nThen test a small acquisition/bonus cohort with unique captain IDs and a retention window.\n\nNo sign-up ROI is identifiable in the current files.")
    sl.notes_slide.notes_text_frame.text='\n\n'.join(x[1] for x in actions)+' '+assumptions
    assert len(prs.slides)==4
    pptx=deliverables/'part_b_presentation.pptx'; prs.save(pptx)
    # Reuse the deterministic PDF renderer used for the Part A presentation.
    from scripts.build_reports import _export_slide_pdf
    _export_slide_pdf(prs, deliverables/'part_b_presentation.pdf', font, bold,
                      'Part B — Airport supply presentation')
    return [memo, deliverables/'part_b_presentation.pdf', pptx, deliverables/'part_b_memo.md']
