"""Build the submission-level two-page memo and six-slide deck for Parts A+B."""
from pathlib import Path
from xml.sax.saxutils import escape
import json


def create_combined_reports(root):
    root = Path(root)
    out = root / 'deliverables'
    out.mkdir(exist_ok=True)
    pa = json.loads((root / 'outputs' / 'part_a' / 'summary.json').read_text())
    pb = json.loads((root / 'outputs' / 'part_b' / 'summary.json').read_text())
    figs_a = root / 'outputs' / 'part_a'
    figs_b = root / 'outputs' / 'part_b'
    decision = 'Run two measured operations pilots—one for onboarding completion and one for airport night supply—before scaling WA002 or airport acquisition.'

    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    font_root = Path('/usr/share/fonts/truetype/dejavu')
    font, bold = 'Helvetica', 'Helvetica-Bold'
    if (font_root / 'DejaVuSans.ttf').exists():
        pdfmetrics.registerFont(TTFont('CombinedSans', str(font_root / 'DejaVuSans.ttf')))
        pdfmetrics.registerFont(TTFont('CombinedSansBold', str(font_root / 'DejaVuSans-Bold.ttf')))
        font, bold = 'CombinedSans', 'CombinedSansBold'
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name='CTitle',fontName=bold,fontSize=18,leading=22,textColor=colors.HexColor('#193B50'),spaceAfter=8))
    styles.add(ParagraphStyle(name='CBody',fontName=font,fontSize=8.55,leading=11.5,spaceAfter=6))
    styles.add(ParagraphStyle(name='CSmall',fontName=font,fontSize=7.5,leading=9.5,spaceAfter=5))
    styles.add(ParagraphStyle(name='CHead',fontName=bold,fontSize=10.5,leading=13,textColor=colors.HexColor('#245B78'),spaceBefore=3,spaceAfter=3))
    def p(v,style='CBody'): return Paragraph(escape(v),styles[style])
    memo = out/'memo.pdf'
    doc = SimpleDocTemplate(str(memo),pagesize=A4,rightMargin=38,leftMargin=38,topMargin=32,bottomMargin=32,
                            title='Captain supply decision memo',author='Supply analytics')
    story=[p('Captain supply: fix completion, then prove airport economics','CTitle'),
           p('DECISION MEMO  |  FOR THE HEAD OF SUPPLY  |  30 JUNE 2026 EXTRACT','CSmall'),p(decision),Spacer(1,3),
           p('What I would decide today','CHead'),
           p(f"Onboarding: of {pa['cohort_signups']:,} January–May signups with a full 30-day follow-up window, {pa['cohort_approved']:,} ({pa['a2o30']:.1%}) were approved and {pa['cohort_active']:,} ({pa['r2a30']:.1%}) completed a first order. The Registration Certificate step is the largest volume loss ({pa['rc_lost']:,} captains). A further {pa['insurance_no_upload']:,} cleared Fitness but had no Insurance upload by day 30."),
           p(f"Airport: between 21:00 and 04:00, only {pb['critical_fulfilment']:.1%} of terminal requests were fulfilled. That leaves {pb['critical_unfulfilled']:,} requests unfulfilled over May–June, about {pb['critical_unfulfilled_per_day']:.0f} per day and {pb['critical_unfulfilled_share']:.1%} of the terminal shortage. Daytime fulfilment is {pb['daytime_fulfilment']:.1%}."),
           p(f"WA002: the simple recipient gap is {pa['campaign_naive_pp']:.1f} percentage points, but every recipient had already cleared Registration Certificate. After aligning the comparison at that stage and accounting for measured pre-send differences, the observed difference is {pa['campaign_association_pp']:+.1f} points (95% range {pa['campaign_ci_pp'][0]:+.1f} to {pa['campaign_ci_pp'][1]:+.1f}). This is not proof that the message caused the improvement."),
           p(f"Airport trips: suburban trips cancel at {pb['suburban_cancel_rate']:.1%} versus {pb['core_cancel_rate']:.1%} in the city core. Only {pb['suburban_next_fare_rate']:.1%} of suburban trips receive a return fare within 20 minutes versus {pb['core_next_fare_rate']:.1%} in the core. This is a directional sample signal; it has no captain ID, completion timestamp, or earnings."),
           p('What I would do on Monday','CHead'),
           p(f"1. Start an Insurance handoff pilot: contact captains 24 hours after a Fitness pass without an Insurance upload, use a comparison group, and target the {pa['insurance_no_upload']:,}-captain pool. The base scenario is about {pa['insurance_monthly_gain']:.0f} additional approvals per month; image-capture repair is another ~{pa['quality_monthly_gain']:.0f} scenario approvals per month."),
           p(f"2. Run a 28-night airport pilot: roster 20 slots from 21:00–04:00, test return matching and a bounded fallback, and measure fulfilled requests per available captain-hour. The 10% recovery target is about {pb['pilot_extra_fulfilled_per_night']:.0f} extra fulfilled requests per night at an expected variable cost of INR {pb['pilot_expected_nightly_variable_cost']:,.0f} per night."),
           p(f"3. Hold the growth requests until the pilots prove value. Randomize about {pa['trial_total']:,} captains who have cleared the Registration Certificate step for a 4-point WA002 test. Do not fund an airport catchment bonus until the night pilot improves fulfilment and return economics."),
           p(f"Data to fix: {pa['mismatch_summary']} document-summary mismatches, {pa['bad_click_delivery']} clicked-but-undelivered messages, and {pb['sample_exceeds_requests_cells']} terminal-hours where the trip sample exceeds hourly request counts.",'CSmall'),PageBreak(),
           p('What I assumed and what would change my answer','CTitle'),
           p('The onboarding scenarios assume that 50% of the identified stalled captains can be recovered and that recovered captains perform like comparable historical captains. The airport scenario assumes 80% attendance, 0.75 completed requests per effective captain-hour, INR 30 per available hour, and an INR 60 fallback with a nightly cap. Costs exclude fixed engineering, airport fees, acquisition bonuses, and captain earnings. These are planning inputs; gains can be zero.'),
           p('I would expand the pilots only if randomized results show incremental approvals or fulfilments, acceptable first-pass quality and complaints, stronger return-fare economics, and Finance-approved unit economics. A null result, poor quality, or no improvement in return utilization would change the recommendation.'),
           p('The campaign result is an observed, adjusted difference—not a proven causal effect. The airport trip comparison is a sample association, not a marketplace-wide revenue estimate. The hourly airport data do not include flight schedules, queue position, dispatch events, unique captains, completion timestamps, or earnings.'),
           p('Working files: Part-A_analysis.ipynb, Part-B_analysis.ipynb, app.py, utils.py, and README.md. The notebooks regenerate the tables and charts used by the dashboard and this memo.','CSmall')]
    def footer(canvas, document):
        canvas.saveState(); canvas.setFont(font,7); canvas.setFillColor(colors.HexColor('#667680'))
        canvas.drawString(38,19,'SUPPLY ANALYTICS  /  PARTS A+B  /  Scenarios are explicit assumptions')
        canvas.drawRightString(A4[0]-38,19,str(document.page)); canvas.restoreState()
    doc.build(story,onFirstPage=footer,onLaterPages=footer)

    # Build a concise six-slide combined deck from existing charts and reportlab's renderer.
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from PIL import Image
    prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
    navy='193B50'; blue='245B78'; orange='DF803C'; green='26836C'; pale='EDF3F6'; grey='5B6B76'
    def box(sl,x,y,w,h,fill='FFFFFF'):
        sh=sl.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h)); sh.fill.solid();sh.fill.fore_color.rgb=RGBColor.from_string(fill);sh.line.fill.background();return sh
    def text(sl,x,y,w,h,value,size=18,color=navy,bold=False):
        sh=sl.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=sh.text_frame;tf.word_wrap=True;tf.margin_left=Inches(.02);tf.margin_right=Inches(.02)
        for i,line in enumerate(value.split('\n')):
            q=tf.paragraphs[0] if i==0 else tf.add_paragraph();q.text=line;q.font.name='DejaVu Sans';q.font.size=Pt(size);q.font.bold=bold;q.font.color.rgb=RGBColor.from_string(color);q.space_after=Pt(8)
        return sh
    def slide(title,sub):
        sl=prs.slides.add_slide(prs.slide_layouts[6]);box(sl,0,0,13.333,.12,blue);text(sl,.55,.35,12.2,.82,title,25,bold=True);text(sl,.58,1.22,12.1,.42,sub,12,color=grey);idx=len(prs.slides);text(sl,.58,7.12,11.8,.22,'PARTS A+B  |  Extract: 30 Jun 2026, 23:59 IST',9,color=grey);text(sl,12.3,7.10,.5,.28,f'{idx}/6',10,color=grey);return sl
    def pic(sl,path,x,y,w,h):
        with Image.open(path) as im: iw,ih=im.size
        scale=min(w/iw,h/ih);fw,fh=iw*scale,ih*scale;sl.shapes.add_picture(str(path),Inches(x+(w-fw)/2),Inches(y+(h-fh)/2),width=Inches(fw),height=Inches(fh))
    def card(sl,x,y,w,h,title,body):box(sl,x,y,w,h,pale);text(sl,x+.18,y+.13,w-.36,.5,title,19,bold=True);text(sl,x+.18,y+.75,w-.36,h-.9,body,15)
    sl=slide('Fix completion; prove airport economics before buying reach',decision)
    for x,val,label in [(0.6,f"{pa['a2o30']:.1%}",'Approved by day 30'),(4.85,f"{pb['critical_fulfilment']:.1%}",'Airport night fulfilment'),(9.1,f"{pa['campaign_association_pp']:+.1f} pp",'WA002 observed difference')]:box(sl,x,1.85,3.62,1.55,pale);text(sl,x+.18,2.0,3.25,.6,val,31,bold=True);text(sl,x+.18,2.68,3.25,.42,label,15)
    text(sl,.65,3.85,12,2.3,f"1   Onboarding: RC is the largest count leak; Insurance and capture quality are recoverable scenarios.\n2   Airport: 21:00–04:00 is a repeatable shortage, but suburban return fares are weak.\n3   Decision: run measured pilots; hold blanket WA002 and airport acquisition scale-up.",21)
    sl=slide('Onboarding loses volume at RC; later handoffs are more recoverable','January–May signups; everyone included; 30-day window.')
    pic(sl,figs_a/'funnel.png',.55,1.75,12.2,4.75);text(sl,.65,6.52,12,.3,f"{pa['insurance_no_upload']:,} Insurance no-upload pool; base scenario ~{pa['insurance_monthly_gain']:.0f} approvals/month. ERickshaw skips Permit.",13,color=grey)
    sl=slide('Airport shortage is concentrated overnight','May discovery, June validation; based on hourly marketplace counts.')
    pic(sl,figs_b/'airport_period_mismatch.png',.55,1.75,12.2,4.75);text(sl,.65,6.52,12,.3,f"Critical: {pb['critical_unfulfilled_per_day']:.0f} unfulfilled/day and {pb['critical_unfulfilled_share']:.1%} of terminal unfulfilled demand; daytime fulfilment {pb['daytime_fulfilment']:.1%}.",13,color=grey)
    sl=slide('Airport trips show weak return utilization','The trip sample is directional and has no unique captain ID or earnings.')
    pic(sl,figs_b/'airport_trip_outcomes.png',.55,1.8,7.2,4.6);card(sl,7.95,1.95,4.75,4.3,'Suburban trips',f"Cancellation: {pb['suburban_cancel_rate']:.1%} vs {pb['core_cancel_rate']:.1%} core.\n\nReturn fare within 20 min: {pb['suburban_next_fare_rate']:.1%} vs {pb['core_next_fare_rate']:.1%}.\n\nFix utilization before buying more headcount.")
    sl=slide('Both growth ideas need experiments','A measured difference or planning scenario is not a proven effect.')
    card(sl,.6,1.95,5.95,4.5,'WA002',f"Observed difference {pa['campaign_association_pp']:+.1f} pp ({pa['campaign_ci_pp'][0]:+.1f} to {pa['campaign_ci_pp'][1]:+.1f}).\n\n~{pa['trial_total']:,} captains who clear Registration Certificate for a 4-point randomised test.\n\nCurrent flow has only {pa['campaign_reach_ceiling']:.2f}× headroom for unique recipients.")
    card(sl,6.8,1.95,5.95,4.5,'Airport',f"No blanket catchment bonus.\n\n20-slot, 28-night timed pilot; target ~{pb['pilot_extra_fulfilled_per_night']:.0f} extra fulfilments/night.\n\nIllustrative variable cost INR {pb['pilot_expected_nightly_variable_cost']:,.0f}/night.\n\nAcquire only after return economics improve.")
    sl=slide('Monday plan and rollout gates','Owners can start diagnosis now; expansion follows incremental outcomes and quality guardrails.')
    rows=[('ONBOARDING','Instrument Insurance stalls and guided recapture; use randomised comparison groups.'),('AIRPORT OPS','Roster timed availability; test queue/return matching and bounded fallback over 28 nights.'),('GROWTH + ANALYTICS','Run the WA002 comparison after Registration Certificate clearance; assess acquisition only with unique IDs and actual costs.')]
    for i,(head,body) in enumerate(rows):y=1.85+i*1.22;box(sl,.6,y,12.1,1.03,pale);text(sl,.8,y+.14,3.1,.68,head,16,bold=True);text(sl,3.95,y+.14,8.5,.78,body,17)
    text(sl,.7,5.92,12,.8,'Gate both pilots on incremental approvals or fulfilled trips, first orders or return-fare guardrails, complaints, fraud/quality and Finance-approved unit economics.',14,color=grey)
    assert len(prs.slides)==6
    from scripts.build_reports import _export_slide_pdf
    _export_slide_pdf(prs,out/'presentation.pdf',font,bold,
                      'Captain supply presentation — Parts A+B')
    return [memo,out/'presentation.pdf']


if __name__ == '__main__':
    create_combined_reports(Path(__file__).resolve().parents[1])
