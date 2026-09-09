"""Presentation-only exports for Part A; all analytical values come from the notebook."""
from pathlib import Path
from xml.sax.saxutils import escape


def create_reports(s, tables, figures, out, deliverables):
    out = Path(out)
    deliverables = Path(deliverables)
    deliverables.mkdir(parents=True, exist_ok=True)
    ci = s['campaign_ci_pp']
    insurance_cpa = s['insurance_monthly_cost'] / s['insurance_monthly_gain']
    quality_cpa = s['quality_monthly_cost'] / s['quality_monthly_gain']
    brief = (
        f"Approve two measured onboarding pilots; hold the blanket 5x WhatsApp budget request. "
        f"Among {s['cohort_signups']:,} January-May signups, {s['cohort_approved']:,} "
        f"({s['a2o30']:.1%}) were approved within 30 days (A2O), and {s['cohort_active']:,} "
        f"({s['r2a30']:.1%}) completed a first order within the same window (R2A). "
        "June does not yet have complete 30-day follow-up."
    )
    findings = [
        ("The largest volume loss is RC; the best late-stage recovery opportunity is Insurance.",
         f"Registration Certificate (RC) loses {s['rc_lost']:,} captains. Separately, {s['insurance_no_upload']:,} captains "
         "clear Fitness but record no Insurance upload within 30 days. Missing upload events do not reveal "
         "whether the problem is document availability, abandonment or an unlogged error; inspect these journeys."),
        ("Low-tier device failures have a concrete product mechanism.",
         f"There are {s['quality_blocked_total']:,} captains blocked at RC, Fitness or Insurance whose latest "
         "failure is blur, low automated-reading confidence or illegibility on a low-tier device. Target capture and recapture quality; "
         "do not relax document or fraud checks."),
        ("The WhatsApp claim overstates the evidence.",
         f"The raw approval difference is {s['campaign_naive_pp']:.1f} percentage points, but every WA002 recipient "
         "had already cleared RC. Among comparable RC clearers, an early send is associated with "
         f"{s['campaign_association_pp']:+.1f} points (95% sampling interval {ci[0]:+.1f} to {ci[1]:+.1f}). "
         "This is an adjusted association, not proven incremental lift: readiness and the assignment rule are unobserved."),
        ("Five times the distinct recipients is infeasible at this flow.",
         f"WA002 already reaches {s['campaign_coverage']:.1%} of eligible RC clearers. Distinct reach has only "
         f"{s['campaign_reach_ceiling']:.2f}x headroom. More repeat touches or a broader audience would be a different intervention.")
    ]
    recs = [
        ("1. Recover the Insurance upload handoff", 
         "Product + city Ops: show the remaining document immediately after Fitness passes; after 24 hours "
         "without a successful upload, send a deep link and offer a support callback. Interview a sample first.",
         f"Base scenario: {s['insurance_no_upload']:,} / 5 months x 50% additional uploads x "
         f"{s['insurance_downstream']:.3f} approval among comparable uploaders = "
         f"{s['insurance_monthly_gain']:.0f} extra approvals/month. At 25%-75% recovery: "
         f"{s['insurance_monthly_gain']*.5:.0f}-{s['insurance_monthly_gain']*1.5:.0f}/month.",
         f"Illustrative cost: {s['insurance_trigger_count']/5:.0f} observable stall contacts/month x INR 20 = "
         f"INR {s['insurance_monthly_cost']:,.0f}/month, or INR {insurance_cpa:.0f}/scenario approval, plus fixed costs. "
         "Risk: captains may lack valid insurance. Test against a randomized stall holdout; measure extra uploads, "
         "A2O/R2A by signup D30, complaints and invalid documents."),
        ("2. Repair image-quality failures, starting with RC",
         "Document product + verification Ops: pilot guided capture, blur/legibility feedback, resumable uploads "
         "and assisted recapture on low-tier devices. Extend the component to Fitness and Insurance.",
         f"Base scenario: across the three documents, blocked latest quality failures / 5 x 50% additional passes "
         f"x downstream approval among comparable low-tier passers = {s['quality_monthly_gain']:.0f} "
         f"extra approvals/month ({s['quality_monthly_gain']*.5:.0f}-{s['quality_monthly_gain']*1.5:.0f} at 25%-75% recovery).",
         f"Illustrative support cost: {s['quality_monthly_targets']:.0f} affected captain-document cases/month x INR 15 = "
         f"INR {s['quality_monthly_cost']:,.0f}/month, or INR {quality_cpa:.0f}/scenario approval, plus engineering. "
         "Randomize before RC upload; measure A2O/R2A, first-pass rates, review workload, device latency and audited false accepts."),
        ("3. Test WA002 before scaling its budget",
         f"Growth + Analytics: randomize approximately {s['trial_total']:,} RC clearers to early WA002 or a suppressed "
         "holdout; keep other communications equal. The design has 80% power for a 4-point lift at a two-sided 5% level. "
         f"Allow about {s['trial_accrual_months']:.1f} months to recruit at historical flow, then 30 days of follow-up.",
         f"Book no causal gain yet. If the association is causal and transports to the remaining "
         f"{s['campaign_extra_recipients_monthly']:.0f} recipients/month, full reach implies about "
         f"{s['campaign_conditional_monthly_gain']:.0f} extra approvals/month. This is conditional, not a forecast.",
         "At an illustrative INR 1/send, messaging cost is the number of assigned sends in rupees, plus tooling and "
         "holdout opportunity cost. Measure approval and activation within 30 days of RC randomization, "
         "cost per incremental approval and opt-outs. Repeated-touch economics remain untested.")
    ]
    caveat = (
        "Scenario assumptions: historical Jan-May flow; 50% incremental recovery; recovered captains convert like "
        "matched vehicle/device peers. Gains halve if their downstream conversion halves, and zero remains possible. "
        "Costs are editable illustrations, not supplied prices; engineering, overhead and acquisition costs are additional. "
        "The two product pools are disjoint at their first unpassed stage; campaign gains overlap and must not be added."
    )
    quality = (
        f"Data owners must reconcile {s['mismatch_summary']:,} June document-summary mismatches, "
        f"{s['bad_click_delivery']:,} clicked-but-undelivered nudges and {s['bad_order_counts']} D7 order totals exceeding D30. "
        "This analysis reconstructs document stages and first orders from timestamps and does not use the inconsistent fields for those outcomes."
    )
    md = ['# Part A — Captain onboarding decision memo',
          'For the Head of Supply · Extract: 30 June 2026, 23:59 IST', brief,
          '## Findings']
    for title, text in findings:
        md += [f'**{title}** {text}']
    md += ['## Ranked actions']
    for title, what, impact, cost in recs:
        md += [f'### {title}', what, impact, cost]
    md += [caveat, quality, 'Working and definitions: `analysis_partA.ipynb`; supporting tables: `outputs/part_a/`.']
    memo_md = deliverables / 'memo.md'
    memo_md.write_text('\n\n'.join(md), encoding='utf-8')

    from reportlab.lib import colors
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import A4
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak, Table, TableStyle
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    font_root = Path('/usr/share/fonts/truetype/dejavu')
    font = 'Helvetica'
    bold = 'Helvetica-Bold'
    if (font_root / 'DejaVuSans.ttf').exists():
        pdfmetrics.registerFont(TTFont('ReportSans', str(font_root / 'DejaVuSans.ttf')))
        pdfmetrics.registerFont(TTFont('ReportSansBold', str(font_root / 'DejaVuSans-Bold.ttf')))
        pdfmetrics.registerFontFamily('ReportSans', normal='ReportSans', bold='ReportSansBold')
        font, bold = 'ReportSans', 'ReportSansBold'
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name='MemoTitle', fontName=bold, fontSize=19, leading=23,
                              textColor=colors.HexColor('#193B50'), spaceAfter=9))
    styles.add(ParagraphStyle(name='MemoBody', fontName=font, fontSize=9.2, leading=13.1, spaceAfter=8))
    styles.add(ParagraphStyle(name='MemoSmall', fontName=font, fontSize=9, leading=12.3, spaceAfter=7))
    styles.add(ParagraphStyle(name='MemoHead', fontName=bold, fontSize=11, leading=14,
                              textColor=colors.HexColor('#245B78'), spaceBefore=5, spaceAfter=5))
    def p(text, style='MemoBody'):
        return Paragraph(escape(text), styles[style])
    memo_pdf = deliverables / 'memo.pdf'
    doc = SimpleDocTemplate(str(memo_pdf), pagesize=A4, rightMargin=40, leftMargin=40,
                            topMargin=35, bottomMargin=35, title='Part A — Onboarding decision memo',
                            author='Supply analytics')
    flow = [p('Captain onboarding: fix completion before buying more reach', 'MemoTitle'),
            p('DECISION MEMO  |  PART A ONLY  |  30 JUNE 2026 EXTRACT', 'MemoSmall'), p(brief), Spacer(1,5)]
    for title, text in findings:
        flow += [p(title, 'MemoHead'), p(text)]
    flow += [Spacer(1,5),p('What changes the decision', 'MemoHead'),
             p('A randomized result with acceptable cost and verification quality can justify rollout. '
               'The current data support specific pilots and a smaller campaign association; they do not support '
               'a blanket 5x spend forecast.'),p(quality,'MemoSmall'),PageBreak(),
             p('Three actions, ranked by recoverable approvals', 'MemoTitle')]
    for title, what, impact, cost in recs:
        flow += [p(title,'MemoHead'),p(what,'MemoSmall'),p(impact,'MemoSmall'),p(cost,'MemoSmall')]
    flow += [Spacer(1,3),p(caveat,'MemoSmall'),p('Audit trail: analysis_partA.ipynb and outputs/part_a/*.csv. '
             'Primary funnel: January-May signups, all signups in denominator, 30 days from signup; '
             'ERickshaw skips Permit. Campaign: complete 30 days after RC clearance, exposure by RC + 24h.','MemoSmall')]
    def footer(canvas, document):
        canvas.saveState(); canvas.setFont(font,7)
        canvas.setFillColor(colors.HexColor('#667680'))
        canvas.drawString(40,21,'SUPPLY ANALYTICS  /  PART A  /  Scenarios are not proven effects')
        canvas.drawRightString(A4[0]-40,21,str(document.page)); canvas.restoreState()
    doc.build(flow,onFirstPage=footer,onLaterPages=footer)

    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
    from pptx.enum.shapes import MSO_SHAPE
    from PIL import Image
    prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
    navy='193B50'; blue='245B78'; orange='DF803C'; pale='EDF3F6'; grey='5B6B76'
    def box(slide,x,y,w,h,fill=None,line=None):
        sh=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h))
        sh.fill.solid(); sh.fill.fore_color.rgb=RGBColor.from_string(fill or 'FFFFFF')
        sh.line.fill.background() if line is None else None
        return sh
    def text(slide,x,y,w,h,value,size=18,color=navy,bold=False):
        sh=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
        tf=sh.text_frame; tf.word_wrap=True
        tf.margin_left=Inches(.02);tf.margin_right=Inches(.02)
        for i,line in enumerate(value.split('\n')):
            para=tf.paragraphs[0] if i==0 else tf.add_paragraph()
            para.text=line;para.font.name='DejaVu Sans';para.font.size=Pt(size)
            para.font.bold=bold;para.font.color.rgb=RGBColor.from_string(color)
            para.space_after=Pt(8)
        return sh
    def slide(title,sub):
        sl=prs.slides.add_slide(prs.slide_layouts[6])
        box(sl,0,0,13.333,.12,blue)
        text(sl,.55,.35,12.2,.7,title,27,bold=True)
        text(sl,.58,1.05,12.1,.65,sub,13,color=grey)
        idx=len(prs.slides)
        text(sl,.58,7.12,11.8,.22,'PART A ONLY  |  Raw CSV analysis  |  Extract: 30 Jun 2026, 23:59 IST',9,color=grey)
        text(sl,12.3,7.10,.5,.28,f'{idx}/6',10,color=grey)
        return sl
    def picture_fit(sl,path,x,y,w,h):
        with Image.open(path) as im: iw,ih=im.size
        scale=min(w/iw,h/ih);fw,fh=iw*scale,ih*scale
        sl.shapes.add_picture(str(path),Inches(x+(w-fw)/2),Inches(y+(h-fh)/2),width=Inches(fw),height=Inches(fh))
    def notes(sl,value): sl.notes_slide.notes_text_frame.text=value
    def card(sl,x,y,w,h,heading,body):
        box(sl,x,y,w,h,pale)
        text(sl,x+.18,y+.13,w-.36,.6,heading,21,bold=True)
        text(sl,x+.18,y+.87,w-.36,h-.98,body,16)

    sl=slide('Fix completion; test WA002 before scaling',
             'Decision: approve Insurance and image-quality pilots; hold the blanket 5x WA002 budget request.')
    for x,value,label in [( .6,f"{s['a2o30']:.1%}",'Approved by D30'),
                          (4.85,f"{s['r2a30']:.1%}",'First order by D30'),
                          (9.1,f"{s['cohort_signups']:,}",'Mature signups')]:
        box(sl,x,1.95,3.62,1.65,pale);text(sl,x+.18,2.1,3.25,.67,value,34,bold=True)
        text(sl,x+.18,2.85,3.25,.45,label,16)
    text(sl,.65,4.0,12,2.4,
         f"1   Insurance handoff: ~{s['insurance_monthly_gain']:.0f} extra approvals/month in the base scenario.\n"
         f"2   Low-device image repair: ~{s['quality_monthly_gain']:.0f}/month in the base scenario.\n"
         f"3   WA002: {s['campaign_association_pp']:+.1f} pp adjusted association; establish incremental lift in a trial.",22)
    notes(sl,brief+' '+caveat)

    sl=slide('RC is the largest volume leak',
             'Jan-May signup cohort; 30 days of follow-up for everyone. Cumulative denominator: all signups.')
    picture_fit(sl,figures['funnel'],.55,1.8,12.2,4.55)
    text(sl,.65,6.45,12,.48,
         'ERickshaw skips Permit. “Lost” means not passed by D30; absence of an upload is not a diagnosed cause.',14,color=grey)
    notes(sl,'Funnel reconstructed from verification events. Conditional stage rates use applicable prerequisite clearers. '
          'June is withheld from the mature D30 trend. Final-review rejections are separate from document loss.')

    sl=slide('Two repairs, with explicit impact assumptions',
             'Base scenarios assume 50% incremental recovery and the downstream conversion of comparable historical peers.')
    card(sl,.6,1.95,5.95,4.5,'1  Insurance upload rescue',
         f"{s['insurance_no_upload']:,} never-uploaded captains after Fitness.\n"
         f"{s['insurance_no_upload']:,} / 5 x 50% x {s['insurance_downstream']:.3f}\n"
         f"= ~{s['insurance_monthly_gain']:.0f} extra approvals/month.\n"
         f"Contact 24h stalls; test a deep link + callback.\n"
         f"Illustrative variable cost: INR {s['insurance_monthly_cost']:,.0f}/month + fixed costs.")
    card(sl,6.8,1.95,5.95,4.5,'2  Low-device image repair',
         f"{s['quality_blocked_total']:,} latest unresolved quality failures across RC, Fitness, Insurance.\n"
         f"Sum: pool / 5 x 50% x downstream A2O\n"
         f"= ~{s['quality_monthly_gain']:.0f} extra approvals/month.\n"
         f"Guided recapture; preserve verification standards.\n"
         f"Illustrative support cost: INR {s['quality_monthly_cost']:,.0f}/month + engineering.")
    text(sl,.65,6.58,12,.35,'Gains halve if rescued captains convert at half the peer rate. Zero remains possible.',13,color=grey)
    notes(sl,'\n\n'.join(' '.join(r) for r in recs[:2])+' '+caveat)

    sl=slide(f"WA002: {s['campaign_association_pp']:+.1f} points after adjustment",
             f"Deck number: {s['campaign_association_pp']:+.1f} pp adjusted association; 95% sampling interval {ci[0]:+.1f} to {ci[1]:+.1f} pp.")
    picture_fit(sl,figures['campaign_estimates'],.55,1.85,12.2,4.35)
    text(sl,.65,6.38,12,.58,
         'Every recipient had cleared RC. Stage alignment and measured adjustment do not remove unmeasured targeting.',15,color=grey)
    notes(sl,'Main analysis: RC + 24h landmark, sent in prior 24h vs no early send, approval by RC + 30 days; '
          f"n={s['campaign_n']:,}. Logistic overlap weighting with 300 full-refit captain bootstrap replicates. "
          'Cross-fitted logistic and gradient boosting are model sensitivities with a different full-population estimand. '
          'No early terminal decisions were excluded in the primary design. Controls can receive later sends. '
          'Confidence: high in counts, moderate in measured-adjusted association, low in causal or 5x forecasts.')

    sl=slide('Reach has a ceiling; test WA002 before expanding',
             'Different audience, extra touches and more signup flow are different interventions.')
    card(sl,.6,1.95,5.95,4.55,'Reach constraint',
         f"{s['campaign_coverage']:.1%} of eligible RC clearers already reached.\n"
         f"Only {s['campaign_reach_ceiling']:.2f}x unique-reach headroom.\n"
         f"~{s['campaign_extra_recipients_monthly']:.0f} untouched captains/month.\n"
         f"~{s['campaign_conditional_monthly_gain']:.0f} extra approvals/month only IF the association is causal and transportable.\n"
         'No causal gain is booked today.')
    card(sl,6.8,1.95,5.95,4.55,'Randomize at RC clearance',
         f"~{s['trial_total']:,} captains total, 1:1 allocation.\n"
         '80% power for a 4 pp lift; 5% two-sided level.\n'
         f"~{s['trial_accrual_months']:.1f} months accrual + 30 days follow-up.\n"
         'Keep other comms equal; suppress holdout WA002.\n'
         'A2O first, then R2A, cost per extra approval and complaints.')
    notes(sl,' '.join(recs[2])+' Costs and value per approved captain are absent; validate actual economics before scaling. '
          'Trial has an assignment-based RC + 30d outcome; this is a new randomized design, not proof from the historical landmark analysis.')

    sl=slide('Monday actions and rollout gates',
             'Owners can start diagnosis now; expansion follows measured incremental approvals and acceptable quality.')
    rows=[('PRODUCT + OPS','Inspect stalled Insurance journeys; ship a 24h stall queue and randomized rescue pilot.'),
          ('DOCUMENT TEAM','Sample failed low-device captures; test guided recapture, latency and false-accept guardrails.'),
          ('GROWTH + ANALYTICS','Launch the RC-stage WA002 holdout; log eligibility, assignment and actual spend.')]
    for i,(head,body) in enumerate(rows):
        y=1.95+i*1.25
        box(sl,.6,y,12.1,1.08,pale);text(sl,.8,y+.15,3.0,.7,head,17,bold=True)
        text(sl,3.95,y+.15,8.5,.8,body,18)
    text(sl,.7,5.94,12,.92,
         f"DATA OWNERS: reconcile {s['mismatch_summary']} document-summary mismatches, "
         f"{s['bad_click_delivery']} clicked-but-undelivered nudges, and {s['bad_order_counts']} inconsistent order totals.\n"
         'Rollout gate: positive incremental A2O, acceptable R2A/fraud/complaints, and Finance-approved unit economics.',15,color=grey)
    notes(sl,quality+' '+caveat)
    assert len(prs.slides)==6
    deck=deliverables/'presentation.pptx';prs.save(deck)
    presentation_pdf=deliverables/'presentation.pdf'
    _export_slide_pdf(prs,presentation_pdf,font,bold)
    return [memo_pdf,presentation_pdf,deck,memo_md]


def _export_slide_pdf(prs, destination, regular_font, bold_font,
                      title='Part A — Captain onboarding presentation'):
    """Render our simple slide primitives to PDF without Office or browser dependencies.

    The editable PPTX and PDF share the exact slide text, charts, positions and
    colors. The supported primitives are rectangles, text boxes and PNG charts.
    """
    from io import BytesIO
    from reportlab.pdfgen.canvas import Canvas
    from reportlab.lib import colors
    from reportlab.lib.utils import ImageReader
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.platypus import Paragraph
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    from pptx.enum.dml import MSO_FILL_TYPE
    unit=12700.0  # EMU per PDF point
    width,height=prs.slide_width/unit,prs.slide_height/unit
    canvas=Canvas(str(destination),pagesize=(width,height))
    canvas.setTitle(title)
    canvas.setAuthor('Supply analytics')
    for slide in prs.slides:
        for shape in slide.shapes:
            x,y,w,h=shape.left/unit,shape.top/unit,shape.width/unit,shape.height/unit
            if shape.shape_type==MSO_SHAPE_TYPE.PICTURE:
                canvas.drawImage(ImageReader(BytesIO(shape.image.blob)),x,height-y-h,w,h,mask='auto')
                continue
            if shape.fill.type==MSO_FILL_TYPE.SOLID:
                canvas.setFillColor(colors.HexColor('#'+str(shape.fill.fore_color.rgb)))
                canvas.rect(x,height-y-h,w,h,fill=1,stroke=0)
            if not shape.has_text_frame or not shape.text.strip():
                continue
            frame=shape.text_frame
            available=w-(frame.margin_left+frame.margin_right)/unit
            cursor=height-y-frame.margin_top/unit
            for para in frame.paragraphs:
                if not para.text:
                    continue
                size=para.font.size.pt if para.font.size else 18
                style=ParagraphStyle('slide',fontName=bold_font if para.font.bold else regular_font,
                    fontSize=size,leading=size*1.15,
                    textColor=colors.HexColor('#'+str(para.font.color.rgb)))
                paragraph=Paragraph(escape(para.text),style)
                _,line_height=paragraph.wrap(available,height)
                paragraph.drawOn(canvas,x+frame.margin_left/unit,cursor-line_height)
                cursor-=line_height+(para.space_after.pt if para.space_after else 0)
        canvas.showPage()
    canvas.save()
